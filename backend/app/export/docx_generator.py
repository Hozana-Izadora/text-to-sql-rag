import asyncio
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.core.logging import get_logger
from app.export.formatting import ColumnFormat, build_column_format_map, format_cell
from app.export.storage import export_filename

logger = get_logger(__name__)

_MAX_TABLE_ROWS = 100
_HEADER_COLOR = "1E3A5F"
_ZEBRA_COLOR = "F2F2F2"
_TEXT_FONT = "Calibri"
_TEXT_SIZE = Pt(11)
_SQL_FONT = "Consolas"
_SQL_SIZE = Pt(9)
_TITLE_COLOR = RGBColor(0x1E, 0x3A, 0x5F)
_MUTED_COLOR = RGBColor(0x66, 0x66, 0x66)

_NUMERIC_FORMATS = {ColumnFormat.CURRENCY, ColumnFormat.DECIMAL_NUMBER, ColumnFormat.INTEGER}
_NARROW_FORMATS = _NUMERIC_FORMATS | {ColumnFormat.DATE, ColumnFormat.DATETIME}


def _shade_cell(cell: Any, hex_color: str) -> None:
    """python-docx não tem API de alto nível pra shading de célula — precisa manipular
    o XML OOXML (w:tcPr/w:shd) diretamente."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_page_number_field(paragraph: Any) -> None:
    """Campo dinâmico de número de página — também sem API de alto nível no python-docx."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _build_header(document: Document) -> None:
    """Logo placeholder: retângulo (tabela 1x1 sombreada) com 'SeguraPro'."""
    header = document.sections[0].header
    table = header.add_table(rows=1, cols=1, width=Inches(2))
    cell = table.cell(0, 0)
    _shade_cell(cell, _HEADER_COLOR)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("SeguraPro")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _add_title_section(document: Document) -> None:
    title = document.add_paragraph()
    title_run = title.add_run("Relatório de Dados")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = _TITLE_COLOR

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        f"SeguraPro Corretora · Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = _MUTED_COLOR


def _add_section_title(document: Document, text: str) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = _TITLE_COLOR


def _add_query_section(document: Document, question: str, sql: str) -> None:
    _add_section_title(document, "Consulta")

    question_paragraph = document.add_paragraph()
    question_paragraph.add_run("Pergunta: ").bold = True
    question_paragraph.add_run(question)

    sql_label = document.add_paragraph()
    sql_label.add_run("SQL executado:").bold = True

    sql_paragraph = document.add_paragraph()
    sql_run = sql_paragraph.add_run(sql)
    sql_run.font.name = _SQL_FONT
    sql_run.font.size = _SQL_SIZE


def _add_results_table(
    document: Document, columns: list[str], rows: list[dict], column_formats: dict[str, ColumnFormat]
) -> None:
    truncated_rows = rows[:_MAX_TABLE_ROWS]

    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_cells = table.rows[0].cells
    for index, column in enumerate(columns):
        header_cells[index].text = column
        _shade_cell(header_cells[index], _HEADER_COLOR)
        header_run = header_cells[index].paragraphs[0].runs[0]
        header_run.bold = True
        header_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Larguras proporcionais: colunas numéricas/data mais estreitas que texto livre.
        width = Inches(1.0) if column_formats.get(column) in _NARROW_FORMATS else Inches(1.6)
        table.columns[index].width = width
        header_cells[index].width = width

    for row_index, row in enumerate(truncated_rows):
        cells = table.add_row().cells
        for col_index, column in enumerate(columns):
            column_format = column_formats.get(column, ColumnFormat.TEXT)
            cells[col_index].text = format_cell(row.get(column), column_format)
            if column_format in _NUMERIC_FORMATS:
                cells[col_index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if row_index % 2 == 1:
                _shade_cell(cells[col_index], _ZEBRA_COLOR)
            cells[col_index].width = table.columns[col_index].width

    if len(rows) > _MAX_TABLE_ROWS:
        note = document.add_paragraph()
        note_run = note.add_run(
            f"Resultados truncados — exibindo os primeiros {_MAX_TABLE_ROWS} de {len(rows)} registros."
        )
        note_run.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = _MUTED_COLOR


def _build_footer(document: Document) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"Gerado por Inquiro em {datetime.now().strftime('%d/%m/%Y')} — Página ")
    run.font.size = Pt(9)
    run.font.color.rgb = _MUTED_COLOR
    _add_page_number_field(paragraph)


def _generate_docx_sync(
    question: str,
    response_text: str,
    sql: str,
    columns: list[str],
    rows: list[dict],
    metadata: dict[str, Any],
    column_formats: dict[str, ColumnFormat],
) -> Path:
    document = Document()
    document.styles["Normal"].font.name = _TEXT_FONT
    document.styles["Normal"].font.size = _TEXT_SIZE

    _build_header(document)
    _add_title_section(document)
    _add_query_section(document, question, sql)

    _add_section_title(document, "Resultados")
    document.add_paragraph(response_text)

    if rows:
        _add_results_table(document, columns, rows, column_formats)

    _build_footer(document)

    output_path = export_filename("docx")
    document.save(output_path)
    return output_path


async def generate_docx(
    question: str,
    response_text: str,
    sql: str,
    columns: list[str],
    rows: list[dict],
    metadata: dict[str, Any],
    column_formats: dict[str, ColumnFormat] | None = None,
) -> Path:
    """Gera um relatório .docx com cabeçalho, seção de consulta, tabela formatada e
    rodapé. python-docx é 100% síncrono — roda em thread separada."""
    resolved_formats = column_formats or build_column_format_map(columns, rows)
    return await asyncio.to_thread(
        _generate_docx_sync, question, response_text, sql, columns, rows, metadata, resolved_formats
    )

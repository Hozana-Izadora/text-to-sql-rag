from decimal import Decimal

import pytest
from docx import Document

from app.export.docx_generator import generate_docx


@pytest.mark.asyncio
async def test_generate_docx_creates_valid_file_with_expected_content() -> None:
    columns = ["broker_name", "total_premium"]
    rows = [
        {"broker_name": "Ana", "total_premium": Decimal("1200.50")},
        {"broker_name": "Bruno", "total_premium": Decimal("800.00")},
    ]

    output_path = await generate_docx(
        question="Qual o total de prêmios por corretor?",
        response_text="Ana teve o maior total de prêmios.",
        sql="SELECT broker_name, total_premium FROM x",
        columns=columns,
        rows=rows,
        metadata={"row_count": 2, "truncated": False},
    )

    try:
        assert output_path.exists()
        assert output_path.suffix == ".docx"

        document = Document(str(output_path))
        full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        assert "Qual o total de prêmios por corretor?" in full_text
        assert "SELECT broker_name, total_premium FROM x" in full_text

        assert len(document.tables) == 1
        table = document.tables[0]
        assert [cell.text for cell in table.rows[0].cells] == columns
        assert len(table.rows) == 1 + len(rows)
    finally:
        output_path.unlink(missing_ok=True)


@pytest.mark.asyncio
async def test_generate_docx_truncates_at_100_rows() -> None:
    columns = ["id"]
    rows = [{"id": i} for i in range(150)]

    output_path = await generate_docx(
        question="pergunta",
        response_text="resposta",
        sql="SELECT id FROM x",
        columns=columns,
        rows=rows,
        metadata={},
    )

    try:
        document = Document(str(output_path))
        table = document.tables[0]
        assert len(table.rows) == 1 + 100

        full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "Resultados truncados" in full_text
    finally:
        output_path.unlink(missing_ok=True)
